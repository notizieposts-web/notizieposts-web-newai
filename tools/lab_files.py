"""Genera i file scaricabili delle sfide "Estrai" in <repo>/lab/ e scrive lab_manifest.json. Uso: python3 lab_files.py <repo>"""
import sys, os, json, random, asyncio, email.utils, datetime as dt
from email.message import EmailMessage
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lab_gen import CH

root = sys.argv[1]; out = os.path.join(root, 'lab'); os.makedirs(out, exist_ok=True)
EXT = {'pdf': 'pdf', 'xlsx': 'xlsx', 'docx': 'docx', 'pptx': 'pptx', 'eml': 'eml', 'png': 'png', 'vtt': 'vtt'}
PDFCSS = '''<style>body{margin:0;font:15px/1.6 "DejaVu Serif",Georgia,serif;color:#121A2B}.paper{padding:28px 34px}
.dh{margin-bottom:10px}.dr{text-align:right;margin-bottom:10px;padding-top:10px;border-top:1px solid #D3D9E6}table{width:100%;border-collapse:collapse;margin:14px 0;font-size:14px}
th,td{border-bottom:1px solid #D3D9E6;padding:4px 8px 4px 4px;text-align:left}td:nth-child(n+2),th:nth-child(n+2){text-align:right}.tot{text-align:right;margin:12px 0}p{margin:8px 0}.ctr{text-align:center}.small{font-size:12px}
.mail{font-family:"DejaVu Sans Mono",monospace;font-size:13px;border-bottom:1px solid #D3D9E6;padding-bottom:8px}.hid{color:#FFFFFF}</style>'''

def pdfs(items):
    from playwright.async_api import async_playwright
    async def run():
        async with async_playwright() as p:
            b = await p.chromium.launch(); pg = await b.new_page()
            for c, path in items:
                await pg.set_content('<!doctype html><html lang="it"><head><meta charset="utf-8">%s</head><body><div class="paper">%s</div></body></html>' % (PDFCSS, c['doc']))
                await pg.pdf(path=path, format='A4', print_background=True, margin={'top': '12mm', 'bottom': '12mm', 'left': '10mm', 'right': '10mm'})
            await b.close()
    asyncio.run(run())

def xlsx(f, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); wb.remove(wb.active)
    for name, rows, hidden in f['sheets']:
        ws = wb.create_sheet(name[:31])
        for row in rows: ws.append(row)
        for cell in ws[1]: cell.font = Font(bold=True)
        if hidden: ws.sheet_state = 'hidden'
    wb.save(path)

def docx(f, path):
    from docx import Document
    d = Document()
    for k, p in enumerate(f['paras']):
        para = d.add_paragraph(p)
        if k == 2: para.runs[0].bold = True
    d.save(path)

def pptx(f, path):
    from pptx import Presentation
    pr = Presentation()
    for t, b in f['slides']:
        s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = t; s.placeholders[1].text = b
    pr.save(path)

def eml(f, path):
    m = EmailMessage(); m['From'] = f['from']; m['To'] = 'clienti@radar-esempio.it'; m['Subject'] = f['subject']
    m['Date'] = email.utils.format_datetime(f['date'].replace(tzinfo=dt.timezone(dt.timedelta(hours=2)))); m.set_content(f['body'])
    open(path, 'wb').write(bytes(m))

def png(f, path, seed):
    from PIL import Image, ImageDraw, ImageFont, ImageFilter
    r = random.Random(seed); W = 520; H = 70 + 34 * len(f['lines'])
    im = Image.new('L', (W, H), 246); dr = ImageDraw.Draw(im)
    fb = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf', 20); fr = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf', 18)
    y = 30
    for k, line in enumerate(f['lines']):
        font = fb if k == 0 or 'TOTALE' in line else fr
        w = dr.textlength(line, font=font); x = (W - w) / 2 if k < 6 else 34
        dr.text((x, y), line, fill=r.randint(20, 60), font=font); y += 34
    for _ in range(1400): dr.point((r.randrange(W), r.randrange(H)), fill=r.randint(150, 230))
    im = im.filter(ImageFilter.GaussianBlur(0.6)).rotate(f['rot'], expand=True, fillcolor=210)
    im.convert('RGB').save(path, optimize=True)

man = {}; pdf_items = []
for c in CH:
    if c['mode'] != 'estrai': continue
    f = c.get('_file') or {'kind': 'pdf'}; ext = EXT[f['kind']]; name = 's%d.%s' % (c['id'], ext); path = os.path.join(out, name)
    if f['kind'] == 'pdf': pdf_items.append((c, path))
    elif f['kind'] == 'xlsx': xlsx(f, path)
    elif f['kind'] == 'docx': docx(f, path)
    elif f['kind'] == 'pptx': pptx(f, path)
    elif f['kind'] == 'eml': eml(f, path)
    elif f['kind'] == 'png': png(f, path, c['id'])
    elif f['kind'] == 'vtt': open(path, 'w', encoding='utf-8').write(f['text'])
    man[c['id']] = name
pdfs(pdf_items)
man = {k: {'f': 'lab/' + v, 'kb': max(1, round(os.path.getsize(os.path.join(out, v)) / 1024))} for k, v in man.items()}
json.dump(man, open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'lab_manifest.json'), 'w'))
print('file', len(man), 'MB', round(sum(os.path.getsize(os.path.join(out, x)) for x in os.listdir(out)) / 1e6, 1))
